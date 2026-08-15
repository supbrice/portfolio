from __future__ import annotations

import json
import os
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from bs4 import BeautifulSoup
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT_DIR = Path(__file__).resolve().parent
DB_PATH = ROOT_DIR / "applications.db"
UPLOADS_DIR = ROOT_DIR / "uploads"
OUTPUT_DIR = ROOT_DIR / "generated"
CLAUDE_API_KEY = os.getenv("ANTHROPIC_API_KEY") or os.getenv("CLAUDE_API_KEY")
ALLOWED_ORIGINS = [
    "http://127.0.0.1:3000",
    "http://localhost:3000",
    "http://127.0.0.1:8000",
    "http://localhost:8000",
    "http://127.0.0.1:5500",
    "http://localhost:5500",
]
VALID_STATUSES = ["drafted", "ready_to_submit", "applied", "interview", "rejected"]

app = FastAPI(title="Portfolio Application Backend")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def init_db() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS submissions (
                id TEXT PRIMARY KEY,
                job_title TEXT NOT NULL,
                company_name TEXT NOT NULL,
                job_link TEXT,
                note TEXT,
                resume_name TEXT,
                status TEXT NOT NULL DEFAULT 'ready_to_submit',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                match_score INTEGER,
                ats_platform TEXT,
                package_path TEXT,
                UNIQUE(company_name, job_title)
            )
            """
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_submissions_company_job ON submissions(company_name, job_title)"
        )


def normalize_status(value: Optional[str]) -> str:
    candidate = (value or "").strip().lower().replace("-", "_").replace(" ", "_")
    if candidate in VALID_STATUSES:
        return candidate
    aliases = {
        "draft": "drafted",
        "ready_to_submit": "ready_to_submit",
        "ready to submit": "ready_to_submit",
        "ready-to-submit": "ready_to_submit",
        "submitted": "applied",
        "interviewing": "interview",
        "accepted": "applied",
        "rejected": "rejected",
    }
    return aliases.get(candidate, "drafted")


def format_status(status: Optional[str]) -> str:
    normalized = normalize_status(status)
    label = normalized.replace("_", " ")
    replacements = {
        "ready to submit": "Ready to Submit",
        "drafted": "Drafted",
        "applied": "Applied",
        "interview": "Interview",
        "rejected": "Rejected",
    }
    return replacements.get(label.lower(), label.title())


def row_to_dict(row: sqlite3.Row) -> Dict[str, Any]:
    row_data = {key: row[key] for key in row.keys()}
    row_data["status"] = normalize_status(row_data.get("status"))
    return row_data


def fetch_submissions() -> List[Dict[str, Any]]:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT * FROM submissions ORDER BY created_at DESC"
        ).fetchall()
        return [row_to_dict(row) for row in rows]


def duplicate_exists(company_name: str, job_title: str) -> bool:
    with sqlite3.connect(DB_PATH) as conn:
        result = conn.execute(
            "SELECT 1 FROM submissions WHERE LOWER(company_name)=LOWER(?) AND LOWER(job_title)=LOWER(?) LIMIT 1",
            (company_name, job_title),
        ).fetchone()
        return result is not None


def save_submission_record(payload: Dict[str, Any]) -> Dict[str, Any]:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        now = utc_now()
        record = {
            "id": payload["id"],
            "job_title": payload["job_title"],
            "company_name": payload["company_name"],
            "job_link": payload.get("job_link") or "",
            "note": payload.get("note") or "",
            "resume_name": payload.get("resume_name") or "",
            "status": payload.get("status") or "ready_to_submit",
            "created_at": payload.get("created_at") or now,
            "updated_at": now,
            "match_score": payload.get("match_score"),
            "ats_platform": payload.get("ats_platform") or "unknown",
            "package_path": payload.get("package_path") or "",
        }
        conn.execute(
            """
            INSERT INTO submissions (
                id, job_title, company_name, job_link, note, resume_name,
                status, created_at, updated_at, match_score, ats_platform, package_path
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                record["id"],
                record["job_title"],
                record["company_name"],
                record["job_link"],
                record["note"],
                record["resume_name"],
                record["status"],
                record["created_at"],
                record["updated_at"],
                record["match_score"],
                record["ats_platform"],
                record["package_path"],
            ),
        )
        return record


def detect_ats_platform(job_link: Optional[str], text: str) -> str:
    link_text = (job_link or "").lower()
    lowered = text.lower()
    if "greenhouse" in link_text or "greenhouse" in lowered:
        return "Greenhouse"
    if "lever" in link_text or "lever" in lowered:
        return "Lever"
    if "workday" in link_text or "workday" in lowered:
        return "Workday"
    if "ashby" in link_text or "ashby" in lowered:
        return "Ashby"
    return "Unknown"


def slugify(value: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9]+", "-", (value or "application")).strip("-")
    return (value or "application").lower()[:40] or "application"


def extract_text_from_resume(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                pages.append(text)
            return "\n".join(pages)
        if suffix == ".docx":
            document = Document(path)
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def extract_skill_terms(text: str, limit: int = 15) -> List[str]:
    keywords = [
        "azure", "aws", "gcp", "terraform", "bicep", "security", "iam", "linux", "python",
        "sql", "power bi", "kubernetes", "docker", "ci/cd", "github actions", "networking",
        "cybersecurity", "incident response", "cloud", "devops", "python", "java", "go",
        "javascript", "typescript", "node", "react", "django", "flask", "automation"
    ]
    found: List[str] = []
    lowered = text.lower()
    for keyword in keywords:
        if keyword in lowered and keyword not in found:
            found.append(keyword)
        if len(found) >= limit:
            break
    if not found:
        pieces = re.findall(r"\b[A-Za-z][A-Za-z+/.-]{2,}\b", text)
        found = sorted(set(piece.lower() for piece in pieces if len(piece) > 3))[:limit]
    return found


def extract_job_keywords(text: str, limit: int = 8) -> List[str]:
    cleaned = re.sub(r"\s+", " ", text)
    phrase_candidates = re.findall(r"(?:[A-Za-z]+(?:[ /-][A-Za-z]+){0,4}){3,}", cleaned)
    preferred = [
        "azure", "aws", "cloud", "security", "networking", "terraform", "python",
        "kubernetes", "linux", "devops", "iam", "incident response", "risk", "governance"
    ]
    hits = [term for term in preferred if term.lower() in cleaned.lower()]
    for candidate in phrase_candidates:
        if len(candidate) > 3 and candidate.lower() not in hits:
            hits.append(candidate.lower())
        if len(hits) >= limit:
            break
    return hits[:limit]


async def call_claude(prompt: str, system_prompt: str) -> Optional[Dict[str, Any]]:
    if not CLAUDE_API_KEY:
        return None
    url = "https://api.anthropic.com/v1/messages"
    headers = {
        "x-api-key": CLAUDE_API_KEY,
        "anthropic-version": "2023-06-01",
        "Content-Type": "application/json",
    }
    payload = {
        "model": "claude-3-5-sonnet-20241022",
        "max_tokens": 1200,
        "messages": [{"role": "user", "content": prompt}],
        "system": system_prompt,
    }
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.post(url, headers=headers, json=payload)
            response.raise_for_status()
            data = response.json()
    except Exception:
        return None

    content_blocks = data.get("content", [])
    text = "".join(block.get("text", "") for block in content_blocks if isinstance(block, dict))
    if not text:
        return None
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        try:
            start = text.find("{")
            end = text.rfind("}")
            if start >= 0 and end > start:
                return json.loads(text[start : end + 1])
        except json.JSONDecodeError:
            return None
        return None


async def analyze_matching(job_title: str, company_name: str, job_link: str, job_description: str, resume_text: str) -> Dict[str, Any]:
    resume_keywords = extract_skill_terms(resume_text)
    job_keywords = extract_job_keywords(job_description)
    overlap = sorted(set(keyword.lower() for keyword in resume_keywords) & set(keyword.lower() for keyword in job_keywords))
    score = min(98, max(15, int((len(overlap) / max(1, len(job_keywords))) * 100))) if job_keywords else 60

    if CLAUDE_API_KEY:
        prompt = (
            "Compare the candidate resume to the job description and return valid JSON with keys "
            "score, rationale, key_requirements, and ats_platform. "
            "The score must be an integer 0-100. Provide 5 concise key requirements.\n\n"
            f"Job title: {job_title}\nCompany: {company_name}\nJob link: {job_link}\n\n"
            f"Resume:\n{resume_text[:5000]}\n\nJob description:\n{job_description[:8000]}"
        )
        system_prompt = "You are an expert recruiting analyst. Use concise but specific matching language. Return only JSON."
        claude_result = await call_claude(prompt, system_prompt)
        if claude_result:
            score = int(claude_result.get("score", score))
            requirements = claude_result.get("key_requirements") or job_keywords
            ats_platform = claude_result.get("ats_platform") or detect_ats_platform(job_link, job_description)
            reasoning = claude_result.get("rationale") or "Review completed."
            return {
                "score": max(0, min(100, score)),
                "reasoning": reasoning,
                "key_requirements": requirements[:5],
                "ats_platform": ats_platform,
            }

    return {
        "score": score,
        "reasoning": "Heuristic fit analysis based on resume and posting keyword overlap.",
        "key_requirements": job_keywords[:5],
        "ats_platform": detect_ats_platform(job_link, job_description),
    }


async def fetch_job_posting(job_title: str, company_name: str, job_link: str) -> Dict[str, Any]:
    fallback = {
        "title": job_title or "Target role",
        "company": company_name or "Target company",
        "description": "No public job posting was provided; matching was based on the information supplied in the form.",
        "ats_platform": "Manual",
    }
    if not job_link:
        return fallback
    try:
        async with httpx.AsyncClient(timeout=20) as client:
            response = await client.get(job_link, follow_redirects=True)
            response.raise_for_status()
    except Exception:
        return fallback

    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    title = soup.title.get_text(" ", strip=True) if soup.title else job_title
    description = text[:12000]
    return {
        "title": title,
        "company": company_name,
        "description": description,
        "ats_platform": detect_ats_platform(job_link, description),
    }


async def generate_document_package(job_title: str, company_name: str, job_link: str, note: str, resume_text: str, match_data: Dict[str, Any]) -> str:
    package_dir = OUTPUT_DIR / f"{slugify(company_name)}-{slugify(job_title)}-{uuid.uuid4().hex[:8]}"
    package_dir.mkdir(parents=True, exist_ok=True)

    summary_payload = {
        "resume_summary": f"Cloud and cybersecurity engineer with a strong record of automating infrastructure, improving operational reliability, and supporting secure production environments.",
        "resume_bullets": [
            f"Led and supported deployments across Azure, AWS, and GCP for secure, scalable cloud workloads.",
            f"Automated operational tasks, configuration workflows, and infrastructure processes to improve reliability and reduce downtime.",
            f"Delivered security-aware engineering practices including IAM reviews, least-privilege access, monitoring, and incident support.",
            f"Bridged technical troubleshooting with clear communication for internal teams, customers, and stakeholders.",
        ],
    }

    if CLAUDE_API_KEY:
        prompt = (
            "Write a tailored application package for the target role. Return valid JSON with keys "
            "resume_summary, resume_bullets, and cover_letter. The summary should be 2-3 sentences. "
            "The cover letter should be 3-5 paragraphs and address the hiring manager.\n\n"
            f"Role: {job_title}\nCompany: {company_name}\nJob link: {job_link}\nNote: {note}\n\n"
            f"Resume excerpt:\n{resume_text[:5000]}\n\n"
            f"Key requirements:\n{', '.join(match_data.get('key_requirements', [])[:5])}"
        )
        system_prompt = "You are an expert career coach and technical writer. Write tailored, professional language for a hiring dossier. Return only JSON."
        claude_doc = await call_claude(prompt, system_prompt)
        if claude_doc:
            summary_payload["resume_summary"] = claude_doc.get("resume_summary") or summary_payload["resume_summary"]
            summary_payload["resume_bullets"] = claude_doc.get("resume_bullets") or summary_payload["resume_bullets"]
            cover_letter = claude_doc.get("cover_letter") or ""
        else:
            cover_letter = (
                f"Dear Hiring Manager,\n\nI am excited to apply for the {job_title} position at {company_name}. "
                f"My background in cloud operations, cybersecurity, automation, and infrastructure support aligns well with the goals of this role. "
                "I enjoy building secure, reliable systems and working cross-functionally to improve operations and customer outcomes.\n\n"
                "Throughout my career, I have focused on troubleshooting production issues, improving platform reliability, and automating manual workflows to reduce operational risk. "
                "I have worked with teams to support cloud environments, improve security posture, and ensure services remain resilient and well monitored.\n\n"
                "I would welcome the opportunity to contribute to {company_name} and help drive technical excellence in a fast-moving environment. "
                "Thank you for your time and consideration. I look forward to the opportunity to speak with you.\n\nSincerely,\nBrice"
            )
    else:
        cover_letter = (
            f"Dear Hiring Manager,\n\nI am excited to apply for the {job_title} position at {company_name}. "
            f"My background in cloud operations, cybersecurity, automation, and infrastructure support aligns well with the goals of this role. "
            "I enjoy building secure, reliable systems and working cross-functionally to improve operations and customer outcomes.\n\n"
            "Throughout my career, I have focused on troubleshooting production issues, improving platform reliability, and automating manual workflows to reduce operational risk. "
            "I have worked with teams to support cloud environments, improve security posture, and ensure services remain resilient and well monitored.\n\n"
            "I would welcome the opportunity to contribute to {company_name} and help drive technical excellence in a fast-moving environment. "
            "Thank you for your time and consideration. I look forward to the opportunity to speak with you.\n\nSincerely,\nBrice"
        )

    docx_path = package_dir / "application-package.docx"
    doc = Document()
    doc.add_heading("Tailored Resume Summary", level=1)
    doc.add_paragraph(summary_payload["resume_summary"])
    doc.add_heading("Key Resume Bullets", level=2)
    for bullet in summary_payload["resume_bullets"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_heading("Cover Letter", level=1)
    doc.add_paragraph(cover_letter)
    doc.save(docx_path)

    pdf_path = package_dir / "application-package.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    width, height = letter
    c.setTitle(f"{job_title} application package")
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, height - 72, f"{company_name} | {job_title}")
    c.setFont("Helvetica-Bold", 13)
    c.drawString(72, height - 110, "Resume Summary")
    c.setFont("Helvetica", 11)
    summary_lines = [
        line for paragraph in summary_payload["resume_summary"].split("\n") for line in wrap_text(paragraph, 90)
    ]
    y = height - 132
    for line in summary_lines[:8]:
        c.drawString(72, y, line)
        y -= 16

    c.setFont("Helvetica-Bold", 13)
    c.drawString(72, y - 18, "Key bullets")
    c.setFont("Helvetica", 11)
    y -= 36
    for bullet in summary_payload["resume_bullets"][:5]:
        for line in wrap_text(f"• {bullet}", 90):
            c.drawString(72, y, line)
            y -= 16
        if y < 72:
            c.showPage();
            y = height - 72

    c.showPage()
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "Cover Letter")
    c.setFont("Helvetica", 11)
    y = height - 96
    for line in wrap_text(cover_letter, 100):
        if y < 72:
            c.showPage();
            y = height - 72
        c.drawString(72, y, line)
        y -= 14
    c.save()

    return str(pdf_path)


def wrap_text(text: str, width: int) -> List[str]:
    words = text.split()
    lines: List[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}" if current else word
        if len(candidate) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


@app.get("/health")
def health_check() -> Dict[str, Any]:
    return {"status": "ok"}


@app.get("/api/applications")
def list_applications() -> Dict[str, Any]:
    return {"success": True, "applications": fetch_submissions()}


@app.post("/api/applications/{submission_id}/status")
async def update_application_status(submission_id: str, request: dict) -> JSONResponse:
    status_value = normalize_status(request.get("status") if isinstance(request, dict) else None)
    with sqlite3.connect(DB_PATH) as conn:
        cursor = conn.execute(
            "UPDATE submissions SET status = ?, updated_at = ? WHERE id = ?",
            (status_value, utc_now(), submission_id),
        )
        if cursor.rowcount == 0:
            return JSONResponse({"success": False, "error": "Submission not found."})
    return JSONResponse({"success": True, "status": status_value, "label": format_status(status_value)})


@app.post("/api/apply")
async def submit_application(
    jobTitle: str = Form(...),
    companyName: str = Form(...),
    jobLink: str = Form(""),
    note: str = Form(""),
    resumeName: str = Form(""),
    resume: UploadFile = File(...),
) -> JSONResponse:
    job_title = (jobTitle or "").strip()
    company_name = (companyName or "").strip()
    job_link = (jobLink or "").strip()
    note_text = (note or "").strip()
    resume_name = (resumeName or resume.filename or "resume.pdf").strip()

    if not job_title or not company_name:
        raise HTTPException(status_code=400, detail="Job title and company name are required.")
    if not resume or not resume.filename:
        raise HTTPException(status_code=400, detail="A resume file is required.")

    if duplicate_exists(company_name, job_title):
        return JSONResponse({"success": False, "error": "A duplicate submission already exists for this company and role."})

    resume_bytes = await resume.read()
    if not resume_bytes:
        return JSONResponse({"success": False, "error": "The uploaded resume file is empty."})

    extension = Path(resume.filename).suffix.lower() if resume.filename else ".pdf"
    safe_name = f"{slugify(company_name)}-{slugify(job_title)}-{uuid.uuid4().hex}{extension}"
    resume_path = UPLOADS_DIR / safe_name
    resume_path.write_bytes(resume_bytes)

    resume_text = extract_text_from_resume(resume_path)
    if not resume_text:
        resume_text = resume_name

    job_posting = await fetch_job_posting(job_title, company_name, job_link)
    match_data = await analyze_matching(job_title, company_name, job_link, job_posting.get("description", ""), resume_text)
    package_file = await generate_document_package(
        job_title=job_title,
        company_name=company_name,
        job_link=job_link,
        note=note_text,
        resume_text=resume_text,
        match_data=match_data,
    )

    submission_id = uuid.uuid4().hex
    package_path = package_file
    db_payload = {
        "id": submission_id,
        "job_title": job_title,
        "company_name": company_name,
        "job_link": job_link,
        "note": note_text,
        "resume_name": resume_name,
        "status": "ready_to_submit",
        "created_at": utc_now(),
        "updated_at": utc_now(),
        "match_score": match_data.get("score"),
        "ats_platform": match_data.get("ats_platform") or detect_ats_platform(job_link, job_posting.get("description", "")),
        "package_path": package_path,
    }
    save_submission_record(db_payload)
    return JSONResponse({"success": True, "submissionId": submission_id})


init_db()


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=False)
